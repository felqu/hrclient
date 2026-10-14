"""XLSX download for the full set of saved vacancies."""
from __future__ import annotations

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.integrations.models.tg_models import JobVacancy
from app.pipeline_log import log_event

_HEADERS = [
    "Канал", "ID сообщения", "Ссылка", "Название", "Компания", "Информация о компании",
    "Город", "Формат работы", "Занятость", "График", "Уровень опыта", "Стаж",
    "Зарплата от", "Зарплата до", "Валюта", "Период оплаты", "Информация о зарплате",
    "Дата публикации",
    "Задачи", "Требования", "Желательные навыки", "Технологии", "Бонусы", "Контакты",
    "Хештеги", "Активна", "Исходный текст",
]


def vacancies_xlsx(vacancies: list[JobVacancy]) -> bytes:
    log_event("xlsx.build.start", count=len(vacancies))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Вакансии"
    sheet.append(_HEADERS)
    for vacancy in vacancies:
        source, company, salary = vacancy.source, vacancy.company, vacancy.salary
        values = [
            source.channel_name if source else None,
            source.message_id if source else None,
            str(source.original_url or source.channel_url or "") if source else "",
            vacancy.title,
            company.name if company else None,
            company.model_dump_json() if company else None,
            vacancy.location,
            vacancy.work_format.value,
            vacancy.employment_type.value,
            vacancy.schedule,
            vacancy.experience_level.value,
            vacancy.experience_years,
            salary.min_amount if salary else None,
            salary.max_amount if salary else None,
            salary.currency if salary else None,
            salary.period if salary else None,
            salary.raw_text if salary else None,
            vacancy.published_date,
            "\n".join(vacancy.tasks),
            "\n".join(vacancy.requirements),
            "\n".join(vacancy.nice_to_have),
            ", ".join(vacancy.technologies),
            "\n".join(vacancy.benefits),
            json.dumps([c.model_dump(mode="json") for c in vacancy.contacts], ensure_ascii=False),
            ", ".join(vacancy.hashtags),
            vacancy.is_active,
            vacancy.raw_text,
        ]
        row_number = sheet.max_row + 1
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column)
            if isinstance(value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
                # Telegram text is data, even when it starts with '='.
                cell.data_type = "s"
            else:
                cell.value = value
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(row=row_number, column=18).number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, header in enumerate(sheet[1], start=1):
        header.font = Font(bold=True, color="FFFFFF")
        header.fill = PatternFill("solid", fgColor="4F46E5")
        header.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = (
            60 if column == len(_HEADERS) else 28
        )
    sheet.row_dimensions[1].height = 32
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    log_event("xlsx.build.done", count=len(vacancies), bytes=output.tell())
    return output.getvalue()
