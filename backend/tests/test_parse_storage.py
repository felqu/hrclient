import os
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import duckdb
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import parse_api
from app.integrations.models.tg_models import Company, JobVacancy, Salary, SourceInfo
from app.services.vacancy_store import VacancyStore, vacancy_id


def vacancy(message_id="42", channel="@jobs", **values):
    return JobVacancy(
        **{
            "title": "Python разработчик",
            "raw_text": "Ищем Python разработчика\nУдалённо",
            "source": SourceInfo(
                channel_name=channel,
                message_id=message_id,
                original_url=f"https://t.me/jobs/{message_id}" if message_id else None,
            ),
            "published_date": date(2026, 10, 14),
            "company": Company(name="Компания"),
            "salary": Salary(min_amount=100000, max_amount=200000),
            "technologies": ["Python", "DuckDB"],
            **values,
        }
    )


class StorageTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "nested" / "vacancies.duckdb"
        env = patch.dict(os.environ, {"PARSE_LOG_PATH": str(Path(self.temp.name) / "pipeline.log")})
        env.start()
        self.addCleanup(env.stop)
        self.store = VacancyStore(self.path)

    def test_persists_full_model_after_reopening_and_updates_message(self):
        original = vacancy()
        self.store.save([original])
        self.assertEqual(VacancyStore(self.path).list(), [original])
        updated = vacancy(channel="https://t.me/jobs", title="Senior Python", raw_text="Новый текст")
        self.store.save([updated, vacancy("43")])
        reloaded = VacancyStore(self.path).list()
        self.assertEqual(len(reloaded), 2)
        self.assertIn(updated, reloaded)
        self.assertEqual(vacancy_id(original), vacancy_id(updated))
        with duckdb.connect(str(self.path)) as connection:
            self.assertEqual(
                connection.execute("SELECT title FROM vacancies WHERE message_id = ?", ["42"]).fetchone(),
                ("Senior Python",),
            )

    def test_empty_import_keeps_records_and_channels_have_distinct_ids(self):
        records = [vacancy(), vacancy(channel="other_jobs")]
        self.store.save(records)
        self.store.save([])
        self.assertCountEqual(self.store.list(), records)

    def test_fallback_without_message_id_and_parallel_imports(self):
        self.store.save([vacancy(None), vacancy(None)])
        with ThreadPoolExecutor(max_workers=3) as executor:
            list(executor.map(lambda index: self.store.save([vacancy(str(index))]), range(6)))
        self.assertEqual(len(self.store.list()), 7)


class ParseAPITests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "vacancies.duckdb"
        env = patch.dict(os.environ, {
            "VACANCIES_DB_PATH": str(self.path),
            "PARSE_LOG_PATH": str(Path(self.temp.name) / "pipeline.log"),
        })
        env.start()
        self.addCleanup(env.stop)
        self.client = TestClient(parse_api.app)
        self.addCleanup(self.client.close)

    def gateway(self, records, errors=()):
        class FakeGateway:
            def __init__(self, **kwargs):
                self.last_errors = list(errors)

            async def search(self, **kwargs):
                return records

        return patch.object(parse_api, "TelegramGateway", FakeGateway)

    def test_both_import_routes_save_and_get_reads_database(self):
        with self.gateway([vacancy()]):
            response = self.client.post("/api/v1/parse", json={"chats": ["jobs"]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["count"], 1)
        with self.gateway([vacancy("43")]):
            response = self.client.post(
                "/api/v1/vacancies/import", json={"telegram_chats": ["jobs"], "sources": ["telegram"]}
            )
        self.assertEqual(response.status_code, 200, response.text)
        # New HTTP client still reads persisted data, rather than the last parse result.
        with TestClient(parse_api.app) as client:
            response = client.get("/api/v1/vacancies")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(len(response.json()), 2)
            self.assertEqual(response.json()[0]["company"], "Компания")
            response = client.get("/api/v1/vacancies/parsed")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(len(response.json()), 2)
            self.assertEqual(response.json()[0]["technologies"], ["Python", "DuckDB"])

    def test_failed_or_empty_parse_does_not_remove_saved_collection(self):
        VacancyStore().save([vacancy()])
        for errors in ([], ["jobs: network failure"]):
            with self.gateway([], errors):
                response = self.client.post("/api/v1/parse", json={"chats": ["jobs"]})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["count"], 0)
            self.assertEqual(len(self.client.get("/api/v1/vacancies").json()), 1)

    def test_storage_failure_is_not_reported_as_success(self):
        with self.gateway([vacancy()]), patch.object(VacancyStore, "save", side_effect=OSError()):
            response = self.client.post("/api/v1/parse", json={"chats": ["jobs"]})
        self.assertEqual(response.status_code, 500)
        self.assertIn("DuckDB", response.json()["detail"])

    def test_xlsx_download_contains_saved_data_and_literal_text(self):
        record = vacancy(title='=HYPERLINK("https://example.com")', raw_text="=1+1\x00")
        VacancyStore().save([record])
        response = self.client.get("/api/v1/vacancies/export.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response.headers["content-type"])
        self.assertIn("attachment;", response.headers["content-disposition"])
        workbook = load_workbook(BytesIO(response.content))
        self.addCleanup(workbook.close)
        sheet = workbook["Вакансии"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["D2"].value, record.title)
        self.assertEqual(sheet["D2"].data_type, "s")
        self.assertEqual(sheet["M2"].value, 100000)
        self.assertEqual(sheet["P2"].value, "month")
        self.assertEqual(sheet["R2"].value.date(), record.published_date)
        self.assertEqual(sheet["AA2"].value, "=1+1")
        self.assertEqual(sheet["AA2"].data_type, "s")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:AA2")
        self.assertEqual(VacancyStore().list()[0].raw_text, record.raw_text)

    def test_empty_collection_download_is_valid_workbook_and_buttons_are_served(self):
        response = self.client.get("/api/v1/vacancies/export.xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.max_row, 1)
        for page in ("/index.html", "/parse.html"):
            response = self.client.get(page)
            self.assertEqual(response.status_code, 200)
            self.assertIn('id="export-button"', response.text)
            self.assertIn('src="export.js?v=2"', response.text)
            self.assertLess(response.text.index('src="api.js?v=2"'),
                            response.text.index('src="export.js?v=2"'))

    def test_export_failure_returns_json_error_and_writes_traceback(self):
        with patch.object(parse_api, "vacancies_xlsx", side_effect=ValueError("xlsx test error")):
            response = self.client.get("/api/v1/vacancies/export.xlsx")
        self.assertEqual(response.status_code, 500)
        self.assertIn("pipeline.log", response.json()["detail"])
        logs = (Path(self.temp.name) / "pipeline.log").read_text()
        self.assertIn("export.failed", logs)
        self.assertIn("xlsx test error", logs)

    def test_export_filename_is_available_to_separate_frontend_origin(self):
        response = self.client.get("/api/v1/vacancies/export.xlsx", headers={"Origin": "http://localhost:5173"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("Content-Disposition", response.headers["access-control-expose-headers"])


if __name__ == "__main__":
    unittest.main()
