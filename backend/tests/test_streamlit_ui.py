import os
import unittest
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import streamlit as st
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from app.integrations.models.tg_models import JobVacancy, SourceInfo
from app.services.parse_jobs import ParseJobManager
from app.services.vacancy_store import VacancyStore

APP_PATH = Path(__file__).resolve().parents[2] / "frontend" / "streamlit_app.py"


class StreamlitUITests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "vacancies.duckdb"
        env = patch.dict(os.environ, {
            "VACANCIES_DB_PATH": str(self.path),
            "PARSE_LOG_PATH": str(Path(self.temp.name) / "pipeline.log"),
            "TG_DEFAULT_CHATS": "",
        })
        env.start()
        self.addCleanup(env.stop)
        st.cache_resource.clear()
        self.addCleanup(st.cache_resource.clear)
        self.store = VacancyStore(self.path)
        self.store.save([
            JobVacancy(title="Python Developer", raw_text="Вакансия Python", technologies=["Python"],
                       published_date=date.today(), source=SourceInfo(channel_name="jobs", message_id="1")),
            JobVacancy(title="Java Developer", raw_text="Вакансия Java", technologies=["Java"],
                       published_date=date.today(), source=SourceInfo(channel_name="jobs", message_id="2")),
        ])

    def app(self):
        app = AppTest.from_file(str(APP_PATH)).run(timeout=10)
        self.assertEqual(len(app.exception), 0, str(app.exception))
        return app

    def button(self, app, label):
        return next(button for button in app.button if button.label == label)

    def test_ui_renders_collection_and_filters_by_search(self):
        app = self.app()
        self.assertIn("Вакансии и парсинг", [element.value for element in app.title])
        app.text_input(key="vacancy_search").set_value("Python").run()
        self.assertEqual(len(app.exception), 0, str(app.exception))
        self.assertTrue(any("Показано 1 из 2" in element.value for element in app.caption))
        self.assertEqual(app.dataframe[0].value.iloc[0]["Вакансия"], "Python Developer")

    def test_xlsx_export_contains_only_filtered_collection(self):
        app = self.app()
        app.text_input(key="vacancy_search").set_value("Python").run()
        self.button(app, "Подготовить XLSX").click().run()
        self.assertEqual(len(app.exception), 0, str(app.exception))
        workbook = load_workbook(BytesIO(app.session_state["export_bytes"]))
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.max_row, 2)
        self.assertEqual(workbook.active["D2"].value, "Python Developer")
        app.text_input(key="vacancy_search").set_value("Java").run()
        self.assertNotIn("export_bytes", app.session_state)

    def test_empty_channels_show_validation_error_without_starting_job(self):
        app = self.app()
        self.button(app, "Запустить парсинг").click().run()
        self.assertEqual(len(app.exception), 0)
        self.assertTrue(any("Укажите хотя бы один канал" in error.value for error in app.error))

    def test_progress_is_visible_while_llm_is_running_and_refresh_does_not_restart_job(self):
        manager = ParseJobManager(self.store)
        entered, release = Event(), Event()

        async def complete(*args, **kwargs):
            entered.set()
            release.wait(5)
            return JobVacancy(title="New vacancy", raw_text="Вакансия Python").model_dump_json()

        backend = SimpleNamespace(name="tg_parser", open=AsyncMock(), close=AsyncMock(), fetch=AsyncMock(
            return_value=[SimpleNamespace(id=3, text="Вакансия Python", date=datetime.now(timezone.utc))],
        ))
        llm = SimpleNamespace(model="test-model", complete=complete, aclose=AsyncMock())
        with patch("app.services.parse_jobs.ParseJobManager", return_value=manager), \
                patch("app.integrations.sources.build_backend", return_value=backend), \
                patch("app.integrations.sources.LLMClient", return_value=llm):
            app = self.app()
            try:
                app.sidebar.text_area[0].set_value("jobs")
                self.button(app, "Запустить парсинг").click().run()
                self.assertTrue(entered.wait(3))
                app.run()
                self.assertEqual(len(app.exception), 0, str(app.exception))
                self.assertEqual(len(manager.jobs), 1)
                self.assertFalse(self.button(app, "Остановить").disabled)
                received = next(metric for metric in app.metric if metric.label == "Получено сообщений")
                self.assertEqual(received.value, "1")
            finally:
                release.set()
                for job in manager.jobs.values():
                    self.assertTrue(job.finished.wait(5))
            app.run()
            self.assertEqual(len(app.exception), 0, str(app.exception))
            self.assertTrue(any("Готово. Сохранено вакансий: 1" in item.value for item in app.success))
            self.assertEqual(len(self.store.list()), 3)


if __name__ == "__main__":
    unittest.main()
